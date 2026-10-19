import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from doc_helper import save_and_archive

def apply_table_styles(table, dxa_widths=None):
    """
    Applies exact department styling:
    - Fixed table width: 7900 dxa (5.486 inches) fitting inside the cell
    - No center misalignment (w:jc is NOT center, stays flush with text)
    - Fixed table layout (w:tblLayout type='fixed')
    - Clean single borders on all edges and inside cells (#CCCCCC)
    - Cell margins (padding): top=80, bottom=80, left=80, right=80 dxa
    - tblHeader on row 0 so headers repeat across page breaks
    - cantSplit on every row so rows don't get sliced across page breaks
    - Exact dxa column widths applied to every cell
    """
    tblPr = table._tbl.tblPr
    
    # Remove any existing jc to prevent centering overflow
    for jc in tblPr.xpath('./w:jc'):
        tblPr.remove(jc)
        
    # Set fixed table width to 7900 dxa
    for w in tblPr.xpath('./w:tblW'):
        tblPr.remove(w)
    tblPr.append(parse_xml(r'<w:tblW %s w:w="7900" w:type="dxa"/>' % nsdecls('w')))
    
    # Set fixed table layout
    for l in tblPr.xpath('./w:tblLayout'):
        tblPr.remove(l)
    tblPr.append(parse_xml(r'<w:tblLayout %s w:type="fixed"/>' % nsdecls('w')))
    
    # Clean borders
    for b in tblPr.xpath('./w:tblBorders'):
        tblPr.remove(b)
    tblBorders = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
        </w:tblBorders>
    ''' % nsdecls('w'))
    tblPr.append(tblBorders)
    
    # Cell margins
    for m in tblPr.xpath('./w:tblCellMar'):
        tblPr.remove(m)
    tblCellMar = parse_xml(r'''
        <w:tblCellMar %s>
            <w:top w:w="80" w:type="dxa"/>
            <w:bottom w:w="80" w:type="dxa"/>
            <w:left w:w="80" w:type="dxa"/>
            <w:right w:w="80" w:type="dxa"/>
        </w:tblCellMar>
    ''' % nsdecls('w'))
    tblPr.append(tblCellMar)

    # Row 0: tblHeader + cantSplit
    header_trPr = table.rows[0]._tr.get_or_add_trPr()
    if not header_trPr.xpath('./w:tblHeader'):
        header_trPr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
    if not header_trPr.xpath('./w:cantSplit'):
        header_trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
    
    # All rows: cantSplit + exact cell width in dxa
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        if not trPr.xpath('./w:cantSplit'):
            trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        if dxa_widths:
            for idx, dxa_w in enumerate(dxa_widths):
                if idx < len(row.cells):
                    cell = row.cells[idx]
                    tcPr = cell._tc.get_or_add_tcPr()
                    for w in tcPr.xpath('./w:tcW'):
                        tcPr.remove(w)
                    tcPr.append(parse_xml(r'<w:tcW %s w:w="%s" w:type="dxa"/>' % (nsdecls('w'), str(dxa_w))))

def style_cell(cell, bg_color=None, bold=False, font_size=9.5, font_name="Calibri", text=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    if bg_color:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)
    if text is not None:
        cell.text = text
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        for r in p.runs:
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.bold = bold

def add_heading(cell, text, bold=True, size=12, keep_with_next=True):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep_with_next
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    return p

def add_subheading(cell, text, bold=True, size=11, keep_with_next=True):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep_with_next
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    return p

def add_body_p(cell, text, bold_prefix=None, space_after=3):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(11)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    return p

def add_figure(cell, image_path, caption_text, width_inches=5.4):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(image_path, width=Inches(width_inches))
    
    cap = cell.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_with_next = True
    c_run = cap.add_run(caption_text)
    c_run.bold = True
    c_run.font.name = "Calibri"
    c_run.font.size = Pt(10)
    return p, cap

def generate_task4_document():
    # Base template for Task 2 to 15
    template_path = os.path.abspath("templates/template_standard.docx")
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found at: {template_path}")
        
    doc = docx.Document(template_path)
    
    # -------------------------------------------------------------
    # 1. Academic Header Envelopes (Table 0 and Table 1)
    # -------------------------------------------------------------
    t0 = doc.tables[0]
    t0.rows[0].cells[1].text = "B.E. Semester VII – Computer Engineering"
    style_cell(t0.rows[0].cells[1], bold=False, font_size=11, text="B.E. Semester VII – Computer Engineering")
    t0.rows[1].cells[1].text = "DevOps Lab"
    style_cell(t0.rows[1].cells[1], bold=False, font_size=11, text="DevOps Lab")
    t0.rows[2].cells[1].text = "Dr. Snehal Andhare"
    style_cell(t0.rows[2].cells[1], bold=False, font_size=11, text="Dr. Snehal Andhare")
    
    t1 = doc.tables[1]
    t1.rows[0].cells[1].text = "Mubina"
    style_cell(t1.rows[0].cells[1], bold=False, font_size=11, text="Mubina")
    t1.rows[1].cells[1].text = "23102A0039"
    style_cell(t1.rows[1].cells[1], bold=False, font_size=11, text="23102A0039")
    
    # Table 2: 5-row standard format
    t2 = doc.tables[2]
    
    # Row 0: Task Number
    t2.rows[0].cells[1].text = "4"
    style_cell(t2.rows[0].cells[1], bold=True, font_size=12, text="4")
    
    # Row 1: Task Title
    task_title = "Git and GitHub Repository Initialization: Create a GitHub repository for the Credit Card Service Portal. Add README, .gitignore, folder structure, issue templates and branch naming rules. Commit the initial application skeleton using meaningful commit messages."
    t2.rows[1].cells[1].text = task_title
    style_cell(t2.rows[1].cells[1], bold=True, font_size=11, text=task_title)
    
    # Row 2: Resources / Apparatus Required
    hw_text = "Hardware: Laptop / PC (x86_64), 16 GB RAM, 512 GB SSD, Intel Core i5/i7 or AMD Ryzen 5/7"
    sw_text = "Software: Git 2.44+, GitHub Enterprise / Cloud, Windows 11 PowerShell, Eclipse Adoptium JDK 17 LTS, Apache Maven 3.9.6, VS Code / IntelliJ IDEA"
    t2.rows[2].cells[1].text = hw_text
    style_cell(t2.rows[2].cells[1], font_size=10, text=hw_text)
    t2.rows[2].cells[2].text = sw_text
    style_cell(t2.rows[2].cells[2], font_size=10, text=sw_text)
    
    # Row 3: Program & Output (Core Deliverable Content)
    body_cell = t2.rows[3].cells[1]
    if body_cell.paragraphs:
        body_cell.paragraphs[0].text = ""
        p0 = body_cell.paragraphs[0]
    else:
        p0 = body_cell.add_paragraph()
        
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.keep_with_next = True
    r0 = p0.add_run("Version control philosophy and repository governance objectives")
    r0.bold = True
    r0.font.name = "Calibri"
    r0.font.size = Pt(12)
    
    add_body_p(body_cell, "In modern enterprise DevOps, Distributed Version Control Systems (DVCS) serve as the authoritative single source of truth for all software artifacts, automated infrastructure code, CI/CD pipeline definitions, and quality verification tests. A properly initialized and governed Git repository guarantees full traceability, atomic rollbacks, peer-reviewed collaborations, and non-destructive experimentation.")
    add_body_p(body_cell, "The objective of Task 4 is to establish a robust version control baseline for the Credit Card Service Portal (CCSP). This entails initializing the local Git repository, defining strict .gitignore exclusion rules to prevent build target and credential leakage, creating a comprehensive README.md architectural manual, authoring structured GitHub Issue Templates for defect reporting and feature enhancement, enforcing a Git Flow branching strategy with Conventional Commits, and committing the verified Spring Boot application skeleton.")

    # Section 1: Directory Structure & Repository Layout
    add_heading(body_cell, "1. Repository directory organization and folder architecture", keep_with_next=True)
    add_body_p(body_cell, "To maintain clean separation between application source code, automated test suites, deployment scripts, academic reports, and evidence screenshots, I organized the workspace into a modular folder hierarchy. Table 4.1 formalizes the directory structure.")

    # Table 4.1: Folder Layout (Total = 7900 dxa: 1800, 1600, 3300, 1200)
    table_dir = body_cell.add_table(rows=8, cols=4)
    apply_table_styles(table_dir, dxa_widths=[1800, 1600, 3300, 1200])
    
    headers_dir = ["Directory Path", "Subsystem Scope", "Key Technical Contents & Artifacts", "Git Tracking Policy"]
    for i, h in enumerate(headers_dir):
        style_cell(table_dir.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    dir_data = [
        (".github/ISSUE_TEMPLATE/", "GitHub Governance", "bug_report.md and feature_request.md issue templates.", "Tracked (Core)"),
        ("app/", "Application Core", "Spring Boot 3.2.5 source code, pom.xml, JPA models, controllers, views.", "Tracked (Core)"),
        ("app/target/", "Build Output", "Compiled .class files, Surefire test reports, executable JAR artifact.", "Ignored (.gitignore)"),
        ("docs/", "Academic Reports", "Canonical university laboratory task Word reports (DevOps_Task{N}_*.docx).", "Tracked (Canonical)"),
        ("docs/prev/", "Version Archival", "Timestamped historical revisions of task documents.", "Tracked (Archival)"),
        ("screenshots/", "Verification Evidence", "Cropped PNG evidence, live browser screenshots, and 300 DPI schematics.", "Tracked (Evidence)"),
        ("scripts/", "DevOps Automation", "Python scripts for programmatic document compilation and diagram generation.", "Tracked (Tooling)")
    ]
    for r_idx, row_vals in enumerate(dir_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_dir.rows[r_idx].cells[c_idx], font_size=8.5, text=val)

    # Section 2: .gitignore Architecture
    add_heading(body_cell, "2. .gitignore specification and asset exclusion policy", keep_with_next=True)
    add_body_p(body_cell, "A robust `.gitignore` file is essential to prevent compiled binaries, operating system temporary files, IDE metadata, and local dependencies from cluttering the remote repository or causing merge conflicts. Table 4.2 presents the categories and pattern rules configured in the root `.gitignore` file.")

    # Table 4.2: Gitignore Rules (Total = 7900 dxa: 1600, 2200, 2600, 1500)
    table_gi = body_cell.add_table(rows=7, cols=4)
    apply_table_styles(table_gi, dxa_widths=[1600, 2200, 2600, 1500])
    
    headers_gi = ["Category", "Exclusion Patterns", "Technical Justification", "DevOps Impact"]
    for i, h in enumerate(headers_gi):
        style_cell(table_gi.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    gi_data = [
        ("Build Artifacts", "target/, **/target/, *.class, *.jar", "Compiled bytecode is reproducible via Maven build tool; prevents repo bloat.", "Clean Repositories"),
        ("IDE Configurations", ".idea/, *.iml, .vscode/, .settings/, .project", "Developer-specific editor settings cause cross-platform configuration drift.", "Environment Neutral"),
        ("OS Temporary Files", ".DS_Store, Thumbs.db, $RECYCLE.BIN/", "Desktop thumbnail and trash caches specific to local Windows/macOS file systems.", "Zero Unrelated Files"),
        ("Log & Crash Dumps", "*.log, logs/, hs_err_pid*, replay_pid*", "Runtime logs and JVM crash dumps contain dynamic, transient execution data.", "Prevents Log Noise"),
        ("Python Cache", "__pycache__/, *.py[cod], .venv/, env/", "Bytecode cache and virtual environments from helper scripting execution.", "Script Cleanliness"),
        ("Local Tools & Temp", "tools/, scratch/, tmp/, *.tmp", "Local JDK/Maven binaries and temporary one-off scratch files.", "Lightweight Commits")
    ]
    for r_idx, row_vals in enumerate(gi_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_gi.rows[r_idx].cells[c_idx], font_size=8.5, text=val)

    # Section 3: README Documentation
    add_heading(body_cell, "3. Master README.md architectural specification", keep_with_next=True)
    add_body_p(body_cell, "The root `README.md` acts as the technical onboarding documentation for developers, evaluators, and automated CI agents. It was authored in standard GitHub Flavored Markdown and encompasses six comprehensive sections:")
    
    add_body_p(body_cell, " Displays student name (Mubina), roll number (23102A0039), semester (B.E. VII), subject evaluator (Dr. Snehal Andhare), and assigned project title.", bold_prefix="• Academic Identification: ")
    add_body_p(body_cell, " High-level overview of the 4-tier monolithic architecture, Spring Boot 3.2.5, H2 in-memory persistence, and Actuator telemetry.", bold_prefix="• System Architecture & Tech Stack: ")
    add_body_p(body_cell, " Itemizes the five essential modules: Card Application, Directory Search, Lifecycle Transitions, KPI Dashboard, and Health Probes.", bold_prefix="• Core MVP Capabilities: ")
    add_body_p(body_cell, " ASCII directory tree mapping all key subdirectories across `.github/`, `app/`, `docs/`, `scripts/`, and `screenshots/`.", bold_prefix="• Folder Architecture Map: ")
    add_body_p(body_cell, " Clear terminal commands for compiling source (`mvn clean compile`), running unit tests (`mvn test`), and launching the local server (`mvn spring-boot:run`).", bold_prefix="• Quick Start Developer Guide: ")
    add_body_p(body_cell, " Summary of branch naming rules (`main`, `dev`, `feature/*`), commit standards, and PR quality gates.", bold_prefix="• Contribution & Branching Policy: ")

    # Section 4: GitHub Issue Templates
    add_heading(body_cell, "4. Structured GitHub issue templates architecture", keep_with_next=True)
    add_body_p(body_cell, "To enforce consistency in defect reporting and requirement gathering, I created two YAML-frontmatter Markdown issue templates in `.github/ISSUE_TEMPLATE/`. Table 4.3 details their structure.")

    # Table 4.3: Issue Templates (Total = 7900 dxa: 1600, 1600, 3200, 1500)
    table_it = body_cell.add_table(rows=3, cols=4)
    apply_table_styles(table_it, dxa_widths=[1600, 1600, 3200, 1500])
    
    headers_it = ["Template File", "Issue Type", "Mandatory Sections & Fields", "DevOps Purpose"]
    for i, h in enumerate(headers_it):
        style_cell(table_it.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    it_data = [
        ("bug_report.md", "Defect / Bug Report", "Bug description, affected component checklist, reproduction steps, expected vs actual behavior, environment details.", "Standardized defect triage & Selenium test failure logging."),
        ("feature_request.md", "Feature Enhancement", "Feature summary, target user story (As a/I want/So that), Gherkin acceptance criteria, proposed technical architecture.", "Structured agile backlog grooming and requirement traceability.")
    ]
    for r_idx, row_vals in enumerate(it_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_it.rows[r_idx].cells[c_idx], font_size=8.5, text=val)

    # Section 5: Branching Strategy & Git Flow
    add_heading(body_cell, "5. Git Flow branching strategy and Conventional Commit rules", keep_with_next=True)
    add_body_p(body_cell, "To enable clean parallel development and conflict-free continuous integration throughout our 15 laboratory tasks, I established a formal branching policy in `BRANCHING_POLICY.md`. Table 4.4 details the branch topology and naming rules.")

    # Table 4.4: Branch Topology (Total = 7900 dxa: 1200, 1600, 2000, 1600, 1500)
    table_br = body_cell.add_table(rows=5, cols=5)
    apply_table_styles(table_br, dxa_widths=[1200, 1600, 2000, 1600, 1500])
    
    headers_br = ["Branch Role", "Branch Name Pattern", "Branch Source & Target", "Merge Governance", "CI/CD Trigger"]
    for i, h in enumerate(headers_br):
        style_cell(table_br.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    br_data = [
        ("Production", "main", "Source: dev; Target: Production", "Protected; tagged releases only (v1.0.0); peer review required.", "Docker CD Release"),
        ("Integration", "dev", "Source: main; Target: main", "Pull Request merge from feature branches; zero test failures.", "Automated Jenkins CI"),
        ("Feature", "feature/<feature-name>", "Source: dev; Target: dev", "Feature development (Task 5 & 6); merged via approved PR.", "Local / Pre-merge CI"),
        ("Hotfix", "hotfix/<defect-id>", "Source: main; Target: main & dev", "Urgent production patch; requires dual merge back into dev.", "Ad-hoc Quality Gate")
    ]
    for r_idx, row_vals in enumerate(br_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_br.rows[r_idx].cells[c_idx], font_size=8.5, text=val)

    # Insert Figure 4.1: Branching Workflow
    fig1_path = os.path.abspath("screenshots/task4/fig4_1_git_branching_workflow.png")
    if os.path.exists(fig1_path):
        add_figure(body_cell, fig1_path, 
                   "Figure 4.1: Credit Card Service Portal — Git Flow Branching & Continuous Integration Architecture", width_inches=5.4)

    # Section 6: Live Execution & Initial Commit
    add_heading(body_cell, "6. Live terminal execution, identity configuration, and initial commit", keep_with_next=True)
    add_body_p(body_cell, "I executed the repository initialization live in the terminal. The sequence included configuring academic identity, staging all project assets, and creating the root commit using Conventional Commits. Table 4.5 details the verified terminal execution log:")

    # Table 4.5: Terminal Log (Total = 7900 dxa: 600, 1800, 3000, 2500)
    table_term = body_cell.add_table(rows=7, cols=4)
    apply_table_styles(table_term, dxa_widths=[600, 1800, 3000, 2500])
    
    headers_term = ["Step", "Git Operation", "Terminal Command Executed", "Verified Output / System Response"]
    for i, h in enumerate(headers_term):
        style_cell(table_term.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    term_data = [
        ("1", "Repository Init", "git init -b main", "Initialized empty Git repository in workspace root."),
        ("2", "User Identity", "git config user.name 'Mubina'\ngit config user.email '23102A0039@...'", "Academic identity configured for all subsequent commit authoring."),
        ("3", "Stage Assets", "git add .", "Staged 51 files across app/, docs/, scripts/, screenshots/, and .github/."),
        ("4", "Initial Commit", "git commit -m 'feat: initialize Credit Card Service Portal application skeleton...'", "[main (root-commit) 2aa4bd5] 51 files changed, 3714 insertions(+)."),
        ("5", "Verify Commit Log", "git log -1 --stat", "Commit 2aa4bd5 authored by Mubina <23102A0039@...> verified."),
        ("6", "Branch Baseline", "git branch dev; git branch -a", "Integration branch 'dev' created successfully off 'main'.")
    ]
    for r_idx, row_vals in enumerate(term_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_term.rows[r_idx].cells[c_idx], font_size=8.5, text=val)

    # Row 4: Conclusion Cell (Formal Academic Conclusion)
    t2.rows[4].cells[0].text = "Conclusion:"
    style_cell(t2.rows[4].cells[0], bold=True, font_size=11, text="Conclusion:")
    
    concl_cell = t2.rows[4].cells[1]
    if concl_cell.paragraphs:
        concl_cell.paragraphs[0].text = ""
        pc = concl_cell.paragraphs[0]
    else:
        pc = concl_cell.add_paragraph()
        
    pc.paragraph_format.space_before = Pt(3)
    pc.paragraph_format.space_after = Pt(3)
    pc.paragraph_format.line_spacing = 1.15
    rc = pc.add_run(
        "In Task 4, I successfully established the distributed version control baseline and repository governance infrastructure "
        "for the Credit Card Service Portal (CCSP). I initialized the Git repository on the main branch, configured the verified student identity "
        "(Mubina, Roll No: 23102A0039), and authored a comprehensive .gitignore file to prevent compilation artifacts, IDE metadata, "
        "and temporary files from cluttering the source repository. I authored a professional, academic README.md detailing the 4-tier architecture, "
        "quick-start execution procedures, and module scopes, and created structured GitHub Issue Templates for defect reporting and feature enhancement. "
        "I formulated a rigorous Git Flow branching policy governed by Conventional Commits standards and established both the main and dev branches. "
        "Finally, I staged the entire application skeleton and executed the root commit (51 files, 3,714 insertions), establishing a tamper-proof, "
        "traceable source baseline for Task 5 Feature Development with Branching."
    )
    rc.font.name = "Calibri"
    rc.font.size = Pt(11)

    # -------------------------------------------------------------
    # Save and Archive Output Document
    # -------------------------------------------------------------
    output_filename = "DevOps_Task4_Mubina_23102A0039.docx"
    canonical_path = os.path.abspath(os.path.join("docs", output_filename))
    
    saved_path = save_and_archive(doc, canonical_path)
    print(f"\nTask 4 Document successfully generated at: {saved_path}")
    return saved_path

if __name__ == "__main__":
    generate_task4_document()
